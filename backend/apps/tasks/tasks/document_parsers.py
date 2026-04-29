"""
文档解析模块：支持 PDF、DOCX、Markdown、URL、GitLab、AI对话 格式的内容提取。
"""
import base64
import io
import os
import re
import tempfile
from typing import Optional

import requests


# ─── PDF 解析 ─────────────────────────────────────────────────────────────────

def parse_pdf(file_path: str) -> dict:
    """
    从 PDF 提取文本和图片，图片按页内位置内联插入。
    返回 { text: str, images: list[bytes] }
    """
    import fitz  # PyMuPDF

    text_parts = []
    images = []
    seen_xrefs = set()
    img_index = 1

    doc = fitz.open(file_path)

    # 有 smask 的 xref 单独 extract_image 会全黑，直接跳过（同组中有 smask=0 的副本包含实际内容）
    xrefs_with_smask = set()
    for page in doc:
        for img_info in page.get_images(full=True):
            xref, smask = img_info[0], img_info[1]
            if smask > 0:
                xrefs_with_smask.add(xref)

    for page_num, page in enumerate(doc, start=1):
        page_parts = []  # list of (y, content_str)

        # 文本块
        text_dict = page.get_text("dict")
        for block in text_dict.get("blocks", []):
            if block["type"] != 0:
                continue
            bbox_y = block.get("bbox", (0, 0, 0, 0))[1]
            lines_text = []
            for line in block.get("lines", []):
                spans_text = "".join(span.get("text", "") for span in line.get("spans", []))
                if spans_text.strip():
                    lines_text.append(spans_text)
            if lines_text:
                page_parts.append((bbox_y, "\n".join(lines_text)))

        # 图片：用 get_image_info 拿带位置的 xref
        try:
            img_info_list = page.get_image_info(xrefs=True)
        except Exception:
            img_info_list = []

        for info in img_info_list:
            xref = info.get("xref", 0)
            if not xref or xref in seen_xrefs or xref in xrefs_with_smask:
                continue
            w = info.get("width", 0)
            h = info.get("height", 0)
            if w < 20 or h < 20:
                continue
            seen_xrefs.add(xref)
            try:
                base_image = doc.extract_image(xref)
                images.append(base_image["image"])
                bbox = info.get("bbox", (0, 0, 0, 0))
                page_parts.append((bbox[1], f"[图片{img_index}]"))
                img_index += 1
            except Exception:
                pass

        page_parts.sort(key=lambda x: x[0])
        if page_parts:
            content = "\n".join(p[1] for p in page_parts)
            text_parts.append(f"<!-- 第 {page_num} 页 -->\n{content}")

    doc.close()
    return {"text": "\n\n".join(text_parts), "images": images}


# ─── DOCX 解析 ────────────────────────────────────────────────────────────────

def parse_docx(file_path: str) -> dict:
    """
    从 DOCX 提取文本、表格和图片，图片在原始位置内联插入占位符。
    返回 { text: str, images: list[bytes] }
    """
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(file_path)
    text_parts = []
    images = []
    img_index = 1

    # 建立 rId -> image blob 映射
    rel_map = {}
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            try:
                rel_map[rel.rId] = rel.target_part.blob
            except Exception:
                pass

    for block in doc.element.body:
        tag = block.tag.split("}")[-1] if "}" in block.tag else block.tag
        if tag == "p":
            # 检查段落中是否有内联图片
            para_parts = []
            has_image = False
            for child in block.iter():
                child_tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
                if child_tag == "t" and child.text:
                    para_parts.append(child.text)
                elif child_tag == "blip":
                    rId = child.get(qn("r:embed")) or child.get(qn("r:link"))
                    if rId and rId in rel_map:
                        images.append(rel_map[rId])
                        para_parts.append(f"[图片{img_index}]")
                        img_index += 1
                        has_image = True

            para_text = "".join(para_parts)
            if para_text.strip():
                style_name = ""
                pPr = block.find(qn("w:pPr"))
                if pPr is not None:
                    pStyle = pPr.find(qn("w:pStyle"))
                    if pStyle is not None:
                        style_name = pStyle.get(qn("w:val"), "")
                if "Heading" in style_name or "heading" in style_name:
                    level = re.search(r"\d+", style_name)
                    prefix = "#" * int(level.group()) if level else "##"
                    text_parts.append(f"{prefix} {para_text}")
                else:
                    text_parts.append(para_text)
            elif has_image:
                text_parts.append(para_text)

        elif tag == "tbl":
            rows = []
            for row in block.findall(".//" + qn("w:tr")):
                cells = []
                for cell in row.findall(".//" + qn("w:tc")):
                    cell_text = "".join(t.text or "" for t in cell.iter() if t.tag.endswith("}t"))
                    cells.append(cell_text.strip())
                rows.append(" | ".join(cells))
            if rows:
                header = rows[0]
                separator = " | ".join(["---"] * len(rows[0].split(" | ")))
                table_md = "\n".join([header, separator] + rows[1:])
                text_parts.append(table_md)

    return {"text": "\n\n".join(text_parts), "images": images}


# ─── Markdown 解析 ────────────────────────────────────────────────────────────

def parse_markdown(content: str, base_url: Optional[str] = None) -> dict:
    """
    解析 Markdown 文本，提取文字并下载图片。
    返回 { text: str, images: list[bytes] }
    """
    images = []
    img_pattern = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")

    def fetch_image(url: str) -> Optional[bytes]:
        try:
            if url.startswith("http://") or url.startswith("https://"):
                resp = requests.get(url, timeout=10)
                if resp.status_code == 200:
                    return resp.content
            elif base_url and not url.startswith("/"):
                full_url = base_url.rstrip("/") + "/" + url
                resp = requests.get(full_url, timeout=10)
                if resp.status_code == 200:
                    return resp.content
        except Exception:
            pass
        return None

    for match in img_pattern.finditer(content):
        img_url = match.group(2)
        img_data = fetch_image(img_url)
        if img_data:
            images.append(img_data)

    # 保留原始 Markdown 文本（图片引用作为占位符）
    return {"text": content, "images": images}


def parse_markdown_file(file_path: str) -> dict:
    """从 Markdown 文件解析。"""
    with open(file_path, "r", encoding="utf-8") as f:
        content = f.read()
    return parse_markdown(content)


# ─── URL 解析 ─────────────────────────────────────────────────────────────────

def parse_url(url: str) -> dict:
    """
    获取 URL 内容并转换为 Markdown。
    返回 { text: str, images: list[bytes] }
    """
    try:
        # 优先尝试 Confluence session 抓取
        from apps.tasks.tasks.md_tasks import _try_fetch_confluence
        confluence_text = _try_fetch_confluence(url)
        if confluence_text is not None:
            return {"text": confluence_text, "images": []}

        resp = requests.get(url, timeout=15, headers={"User-Agent": "Mozilla/5.0"})
        resp.raise_for_status()
        content_type = resp.headers.get("content-type", "")

        if "html" in content_type:
            # HTML → Markdown
            try:
                from html2text import html2text
                md_text = html2text(resp.text)
            except ImportError:
                # 降级：简单去除 HTML 标签
                md_text = re.sub(r"<[^>]+>", "", resp.text)
            return parse_markdown(md_text, base_url=url)
        elif "markdown" in content_type or url.endswith(".md"):
            return parse_markdown(resp.text, base_url=url)
        else:
            return {"text": resp.text, "images": []}
    except Exception as e:
        raise ValueError(f"URL 获取失败: {e}")


# ─── GitLab issue 解析 ────────────────────────────────────────────────────────

def parse_gitlab_issue(issue_url: str, gitlab_token: Optional[str] = None) -> dict:
    """
    通过 GitLab API 获取 issue 内容。
    支持格式：https://gitlab.com/group/project/-/issues/123
    返回 { text: str, images: list[bytes] }
    """
    # 解析 URL 提取 project path 和 issue ID
    match = re.match(r"(https?://[^/]+)/(.+?)/-/issues/(\d+)", issue_url)
    if not match:
        raise ValueError(f"无法解析 GitLab issue URL: {issue_url}")

    base_url, project_path, issue_id = match.groups()
    project_path_encoded = project_path.replace("/", "%2F")
    api_url = f"{base_url}/api/v4/projects/{project_path_encoded}/issues/{issue_id}"

    headers = {}
    if gitlab_token:
        headers["PRIVATE-TOKEN"] = gitlab_token

    try:
        resp = requests.get(api_url, headers=headers, timeout=15)
        resp.raise_for_status()
        issue = resp.json()
    except Exception as e:
        raise ValueError(f"GitLab API 请求失败: {e}")

    title = issue.get("title", "")
    description = issue.get("description", "") or ""
    labels = ", ".join(issue.get("labels", []))
    state = issue.get("state", "")

    text = f"# {title}\n\n"
    if labels:
        text += f"**标签**: {labels}  \n"
    text += f"**状态**: {state}\n\n"
    text += description

    # 获取评论
    notes_url = f"{base_url}/api/v4/projects/{project_path_encoded}/issues/{issue_id}/notes"
    try:
        notes_resp = requests.get(notes_url, headers=headers, timeout=15)
        if notes_resp.status_code == 200:
            notes = notes_resp.json()
            for note in notes:
                if not note.get("system", False):
                    author = note.get("author", {}).get("name", "")
                    body = note.get("body", "")
                    text += f"\n\n---\n**评论（{author}）**:\n{body}"
    except Exception:
        pass

    # 下载描述中的图片附件
    result = parse_markdown(text, base_url=base_url)
    return result


# ─── AI 对话解析 ──────────────────────────────────────────────────────────────

def parse_ai_conversation(content: str) -> dict:
    """
    解析 AI 对话记录，提取需求相关内容。
    content 格式：纯文本对话记录
    返回 { text: str, images: list[bytes] }
    """
    lines = content.strip().split("\n")
    structured_parts = ["# AI 对话需求提取\n"]
    current_role = None
    current_block = []

    for line in lines:
        # 检测角色标记
        if re.match(r"^(用户|User|Human|我|提问者)[：:]\s*", line):
            if current_block and current_role:
                structured_parts.append(f"**{current_role}**: " + " ".join(current_block))
            current_role = "用户"
            current_block = [re.sub(r"^(用户|User|Human|我|提问者)[：:]\s*", "", line)]
        elif re.match(r"^(AI|助手|Assistant|Claude|GPT|回答)[：:]\s*", line):
            if current_block and current_role:
                structured_parts.append(f"**{current_role}**: " + " ".join(current_block))
            current_role = "AI"
            current_block = [re.sub(r"^(AI|助手|Assistant|Claude|GPT|回答)[：:]\s*", "", line)]
        elif line.strip():
            current_block.append(line)
        else:
            if current_block and current_role:
                structured_parts.append(f"**{current_role}**: " + " ".join(current_block))
                current_block = []

    if current_block and current_role:
        structured_parts.append(f"**{current_role}**: " + " ".join(current_block))

    return {"text": "\n\n".join(structured_parts), "images": []}


# ─── 图片识别（Vision LLM）────────────────────────────────────────────────────

def recognize_images_with_vision_llm(images: list, batch_size: int = 1) -> list:
    """
    使用 Vision LLM 逐张识别图片内容（默认每次只发1张，避免大请求超时）。
    images: list[bytes]
    返回: list[str]  每张图片的文字描述（纯描述，不含 "[图片 N]:" 前缀）
    """
    from apps.integrations.views import AIService

    descriptions = []
    for i in range(0, len(images), batch_size):
        batch = images[i:i + batch_size]
        temp_paths = []
        try:
            for img_bytes in batch:
                tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
                tmp.write(img_bytes)
                tmp.close()
                temp_paths.append(tmp.name)

            if len(temp_paths) == 1:
                # 单张图片：直接描述，不需要分隔符
                prompt = (
                    "请详细描述这张图片的内容。\n"
                    "如果是 UI 界面截图，描述界面布局、元素和标注信息；\n"
                    "如果是流程图，描述流程步骤和关系；\n"
                    "如果是表格，提取表格内容。\n"
                    "输出结构化文本描述。"
                )
            else:
                # 多张图片：要求分隔符
                prompt = (
                    f"我将提供 {len(batch)} 张图片，请逐张详细描述每张图片的内容。\n\n"
                    "要求：\n"
                    "1. 按顺序描述每张图片，使用 '---图片N---' 作为分隔符（N从1开始）\n"
                    "2. 如果是 UI 界面截图，描述界面布局、元素和标注信息\n"
                    "3. 如果是流程图，描述流程步骤和关系\n"
                    "4. 如果是表格，提取表格内容\n"
                    "5. 每张图片的描述要独立完整\n\n"
                    "输出格式示例：\n"
                    "---图片1---\n"
                    "[第1张图片的详细描述]\n\n"
                    "---图片2---\n"
                    "[第2张图片的详细描述]"
                )

            system = "你是一个专业的 UI/UX 分析师，擅长描述产品原型图和流程图。"
            result = AIService.complete_with_images(prompt, temp_paths, system=system)

            if len(temp_paths) == 1:
                descriptions.append(result.strip())
            else:
                parts = re.split(r'---图片\d+---', result)
                parts = [p.strip() for p in parts if p.strip()]
                for j in range(len(temp_paths)):
                    descriptions.append(parts[j] if j < len(parts) else result)
        finally:
            for path in temp_paths:
                try:
                    os.remove(path)
                except OSError:
                    pass

    return descriptions



# ─── 图片保存 + 占位符工具 ────────────────────────────────────────────────────

def save_images_and_build_text(text: str, images: list, save_dir: str, start_index: int = 1) -> tuple:
    """
    将图片字节列表保存到磁盘。
    text 中已由解析器内联了 [图片N] 占位符，此函数只负责保存文件，不再追加占位符。
    返回：(text: str, saved_image_paths: list[str], next_index: int)
    """
    import uuid
    os.makedirs(save_dir, exist_ok=True)

    saved_paths = []
    idx = start_index

    for img_bytes in images:
        filename = f"img_{idx}_{uuid.uuid4().hex[:8]}.png"
        filepath = os.path.join(save_dir, filename)
        with open(filepath, 'wb') as f:
            f.write(img_bytes)
        saved_paths.append(filepath)
        idx += 1

    return text, saved_paths, idx


def save_images_from_markdown(text: str, images: list, save_dir: str, start_index: int = 1) -> tuple:
    """
    专用于 Markdown：将 ![alt](url) 图片语法替换为 [图片N] 占位符，
    同时将已下载的图片字节保存到磁盘。

    返回：(text_with_placeholders: str, saved_image_paths: list[str], next_index: int)
    """
    import uuid
    os.makedirs(save_dir, exist_ok=True)

    saved_paths = []
    idx = start_index
    img_pattern = re.compile(r'!\[([^\]]*)\]\(([^)]+)\)')

    def replacer(match):
        nonlocal idx
        placeholder = f'[图片{idx}]'
        idx += 1
        return placeholder

    text_with_placeholders = img_pattern.sub(replacer, text)

    # 保存图片字节
    for img_bytes in images:
        filename = f"img_{start_index + len(saved_paths)}_{uuid.uuid4().hex[:8]}.png"
        filepath = os.path.join(save_dir, filename)
        with open(filepath, 'wb') as f:
            f.write(img_bytes)
        saved_paths.append(filepath)

    return text_with_placeholders, saved_paths, idx

def build_parsed_content(text: str, images: list, enable_image_recognition: bool) -> tuple:
    """
    构建两个内容字段：
    - parsed_content: 纯文本 + 表格
    - parsed_content_with_images: 文本 + 表格 + 图片描述（如启用）
    返回 (parsed_content, parsed_content_with_images)
    """
    parsed_content = text.strip()

    if not enable_image_recognition or not images:
        return parsed_content, parsed_content

    image_descriptions = recognize_images_with_vision_llm(images)
    image_section = "\n\n## 图片内容识别\n\n" + "\n\n".join(image_descriptions)
    parsed_content_with_images = parsed_content + image_section

    return parsed_content, parsed_content_with_images
